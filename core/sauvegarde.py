"""Sauvegarde et restauration des données d'un atelier.

La sauvegarde est strictement cloisonnée par atelier (multi-boutique) :
seuls les enregistrements dont le FK `atelier` correspond à l'atelier
courant sont exportés, et une restauration refuse tout fichier dont les
enregistrements viseraient un autre atelier.

Contenu : les douze modèles métier (clients, mensurations, catalogue,
accessoires, commandes et leurs lignes, paiements, mouvements de stock,
dépenses, employés, paies). Les comptes d'accès, l'abonnement et les
fichiers média (photos, logos) ne sont PAS inclus : restaurer des
comptes pourrait verrouiller l'accès à l'atelier.

Les fichiers sont déposés dans <BASE_DIR>/sauvegardes/ sous le nom
sauvegarde_<idAtelier>_<AAAAMMJJ-HHMMSS>.xlsx : un classeur Excel lisible
(une feuille par modèle) contenant en feuille masquée « _RESTAURATION »
la charge JSON nécessaire à une restauration exacte.
"""
import json
import re
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

from django.apps import apps
from django.conf import settings
from django.core import serializers
from django.db import transaction

# Ordre de chargement : parents d'abord ( FK satisfaits ).
MODELES_SAUVEGARDES = [
    'Client', 'CatalogueModele', 'Accessoire', 'Commande', 'Mensuration',
    'Depense', 'MouvementStock', 'Paiement', 'LigneCommande',
    'LigneAccessoire', 'Employe', 'Paie',
]
# Suppression avant restauration : enfants d'abord.
ORDRE_SUPPRESSION = list(reversed(MODELES_SAUVEGARDES))

FEUILLES = {
    'Client': 'Clients',
    'Mensuration': 'Mensurations',
    'CatalogueModele': 'Modèles',
    'Accessoire': 'Accessoires',
    'Commande': 'Commandes',
    'LigneCommande': 'Lignes de commande',
    'LigneAccessoire': 'Lignes accessoires',
    'Paiement': 'Paiements',
    'MouvementStock': 'Mouvements de stock',
    'Depense': 'Dépenses',
    'Employe': 'Employés',
    'Paie': 'Paies',
}
CHARGE_MAX = 30000  # limite de caractères par cellule Excel

MOTIF_NOM = re.compile(r'^sauvegarde_(\d+)_(\d{8}-\d{6})\.xlsx$')


def _modele(nom):
    return apps.get_model('core', nom)


def dossier_sauvegardes():
    dossier = Path(settings.BASE_DIR) / 'sauvegardes'
    dossier.mkdir(parents=True, exist_ok=True)
    return dossier


def chemin_sauvegarde(atelier, nom):
    """Chemin d'une sauvegarde si elle existe et appartient à l'atelier."""
    motif = MOTIF_NOM.match(nom or '')
    if not motif or motif.group(1) != str(atelier.id):
        return None
    chemin = dossier_sauvegardes() / nom
    return chemin if chemin.is_file() else None


def _valeur_lisible(valeur, field, cache):
    """Valeur de champ transformée pour lecture humaine dans Excel."""
    if valeur is None:
        return ''
    if field is not None and field.many_to_many or field.many_to_one or field.one_to_one:
        modele = field.related_model
        if modele is not None:
            table = cache.setdefault(modele, {o.pk: str(o) for o in modele.objects.all()})
            if isinstance(valeur, list):
                return ', '.join(table.get(v, v) for v in valeur)
            return table.get(valeur, valeur)
    if isinstance(valeur, list):
        return ', '.join(str(v) for v in valeur)
    if isinstance(valeur, dict):
        return json.dumps(valeur, ensure_ascii=False)
    if isinstance(valeur, datetime):
        return valeur.replace(tzinfo=None)  # Excel ne gère pas les fuseaux
    if isinstance(valeur, date):
        return valeur
    return valeur


def _ecrire_classeur(modeles_objets, enveloppe_json, infos):
    """Construit le classeur Excel : feuilles lisibles + charge masquée."""
    wb = Workbook()
    wb.remove(wb.active)
    feuille = wb.create_sheet('Informations')
    for cle, val in infos:
        feuille.append([cle, val])
    cache = {}
    for nom, objets in modeles_objets:
        ws = wb.create_sheet(FEUILLES.get(nom, nom)[:31])
        if not objets:
            ws.append(['(aucun enregistrement)'])
            continue
        champs = []
        for objet in objets:
            for cle in objet['fields']:
                if cle not in champs:
                    champs.append(cle)
        modele = _modele(nom)
        entetes = ['Identifiant']
        for c in champs:
            try:
                entetes.append(str(modele._meta.get_field(c).verbose_name).capitalize())
            except Exception:
                entetes.append(c.capitalize().replace('_', ' '))
        ws.append(entetes)
        for objet in objets:
            ligne = [objet['pk']]
            for cle in champs:
                field = None
                try:
                    field = modele._meta.get_field(cle)
                except Exception:
                    field = None
                ligne.append(_valeur_lisible(objet['fields'].get(cle), field, cache))
            ws.append(ligne)
    # Charge de restauration, masquée, découpée en tronçons < 32 767 caractères
    ws = wb.create_sheet('_RESTAURATION')
    ws.sheet_state = 'hidden'
    for i in range(0, len(enveloppe_json), CHARGE_MAX):
        ws.append([enveloppe_json[i:i + CHARGE_MAX]])
    return wb


def creer_sauvegarde(atelier):
    """Exporte les données de l'atelier en Excel ; retourne le nom du fichier."""
    modeles_objets = []
    objets = []
    for nom in MODELES_SAUVEGARDES:
        qs = list(_modele(nom).objects.filter(atelier=atelier).order_by('pk'))
        modeles_objets.append((nom, qs))
        objets.extend(qs)
    contenu = json.loads(serializers.serialize('json', objets))
    enveloppe = {
        'version': 1,
        'application': 'Projet_tailleur',
        'atelier_id': atelier.id,
        'atelier_nom': atelier.nom,
        'cree_le': datetime.now().isoformat(timespec='seconds'),
        'objets': contenu,
    }
    charge = json.dumps(enveloppe, ensure_ascii=False)
    wb = _ecrire_classeur(
        [(nom, json.loads(serializers.serialize('json', qs))) for nom, qs in modeles_objets],
        charge,
        [
            ('Application', 'Projet_tailleur'),
            ('Atelier', atelier.nom),
            ('Identifiant atelier', atelier.id),
            ('Créée le', datetime.now().strftime('%d/%m/%Y %H:%M')),
            ('Enregistrements', len(objets)),
            ('Feuille masquée', '_RESTAURATION (charge de restauration)'),
        ],
    )
    nom = f"sauvegarde_{atelier.id}_{datetime.now():%Y%m%d-%H%M%S}.xlsx"
    wb.save(dossier_sauvegardes() / nom)
    return nom


def lister_sauvegardes(atelier):
    """Sauvegardes de l'atelier, plus récentes d'abord."""
    resultats = []
    for chemin in dossier_sauvegardes().glob('sauvegarde_*.xlsx'):
        motif = MOTIF_NOM.match(chemin.name)
        if not motif or motif.group(1) != str(atelier.id):
            continue
        estampille = datetime.strptime(motif.group(2), '%Y%m%d-%H%M%S')
        resultats.append({
            'nom': chemin.name,
            'date': estampille,
            'taille': chemin.stat().st_size,
        })
    resultats.sort(key=lambda e: e['date'], reverse=True)
    return resultats


def supprimer_sauvegarde(atelier, nom):
    chemin = chemin_sauvegarde(atelier, nom)
    if chemin is None:
        return False
    chemin.unlink()
    return True


def _objets_du_fichier(contenu):
    """Valide l'enveloppe et retourne (objets, erreur)."""
    try:
        enveloppe = json.loads(contenu)
    except (ValueError, UnicodeDecodeError):
        return None, "Le fichier n'est pas une sauvegarde JSON valide."
    if not isinstance(enveloppe, dict) or 'objets' not in enveloppe:
        return None, "Structure de sauvegarde non reconnue."
    objets = enveloppe['objets']
    if not isinstance(objets, list):
        return None, "Structure de sauvegarde non reconnue."
    return objets, None


def objets_pour_atelier(atelier, objets):
    """Vérifie que chaque objet appartient au modèle et à l'atelier."""
    autorises = {f'core.{nom.lower()}' for nom in MODELES_SAUVEGARDES}
    for objet in objets:
        if objet.get('model') not in autorises:
            return f"Modèle inattendu dans la sauvegarde : {objet.get('model')}."
        if objet.get('fields', {}).get('atelier') != atelier.id:
            return ("La sauvegarde contient des données d'un autre atelier : "
                    "restauration refusée.")
    return None


def _charge_depuis_fichier(contenu, nom_fichier=''):
    """Retourne (objets, erreur) depuis un .xlsx (feuille masquée) ou .json."""
    if nom_fichier.lower().endswith('.xlsx') or contenu[:2] == b'PK':
        donnees = contenu if isinstance(contenu, bytes) else contenu.encode('latin-1', 'replace')
        try:
            wb = load_workbook(BytesIO(donnees), read_only=True)
        except Exception:
            return None, "Fichier Excel illisible ou corrompu."
        if '_RESTAURATION' not in wb.sheetnames:
            return None, ("Ce classeur ne contient pas de charge de "
                          "restauration (feuille masquée _RESTAURATION).")
        troncons = []
        for ligne in wb['_RESTAURATION'].iter_rows(values_only=True):
            if ligne and ligne[0]:
                troncons.append(str(ligne[0]))
        wb.close()
        contenu = ''.join(troncons)
    elif isinstance(contenu, bytes):
        contenu = contenu.decode('utf-8', errors='replace')
    return _objets_du_fichier(contenu)


def restaurer_sauvegarde(atelier, contenu, nom_fichier=''):
    """Remplace les données de l'atelier par celles du fichier.

    Retourne (succès, message). Tout se passe dans une transaction :
    en cas d'erreur, rien n'est modifié.
    """
    objets, erreur = _charge_depuis_fichier(contenu, nom_fichier)
    if erreur:
        return False, erreur
    erreur = objets_pour_atelier(atelier, objets)
    if erreur:
        return False, erreur
    try:
        with transaction.atomic():
            for nom in ORDRE_SUPPRESSION:
                _modele(nom).objects.filter(atelier=atelier).delete()
            nombre = 0
            flux = json.dumps(objets)
            for deserialise in serializers.deserialize('json', flux):
                deserialise.save()
                nombre += 1
    except Exception as exc:  # fichier corrompu, contraintes…
        return False, f"Restauration interrompue, aucune donnée modifiée : {exc}"
    return True, f"Restauration terminée : {nombre} enregistrements rétablis."


# ==========================================
# SAUVEGARDE SYSTÈME (super-admin) : classeur Excel global, archivage
# ==========================================
DOSSIER_SYSTEME = 'sauvegardes_systeme'
MOTIF_NOM_SYSTEME = re.compile(r'^sauvegarde_systeme_(\d{8}-\d{6})\.xlsx$')


def dossier_systeme():
    dossier = Path(settings.BASE_DIR) / DOSSIER_SYSTEME
    dossier.mkdir(parents=True, exist_ok=True)
    return dossier


def creer_sauvegarde_systeme(atelier=None):
    """Classeur Excel des données — toute la plateforme (atelier=None) ou
    une seule boutique (atelier fourni). Les modèles communs de la plateforme
    (formules d'abonnement) sont toujours exportés intégralement ; les
    modèles locataires sont filtrés par atelier le cas échéant."""
    from django.contrib.auth.models import User
    from .models import Atelier, Abonnement, PlanAbonnement, Profil

    wb = Workbook()
    wb.remove(wb.active)
    infos = wb.create_sheet('Informations')
    infos.append(['Application', 'Projet_tailleur'])
    infos.append(['Portée', atelier.nom + ' (boutique unique)' if atelier else 'Plateforme complète (toutes boutiques)'])
    infos.append(['Créée le', datetime.now().strftime('%d/%m/%Y %H:%M')])

    def nettoie(valeur):
        if isinstance(valeur, datetime):
            return valeur.replace(tzinfo=None)
        return valeur

    def feuille(nom, entetes, lignes):
        ws = wb.create_sheet(nom[:31])
        ws.append(entetes)
        for ligne in lignes:
            ws.append([nettoie(v) for v in ligne])
        return ws

    for nom in MODELES_SAUVEGARDES:
        qs = _modele(nom).objects.all().order_by('pk')
        if atelier is not None:
            champs_modele = {f.name for f in _modele(nom)._meta.get_fields()}
            if 'atelier' in champs_modele:
                qs = qs.filter(atelier=atelier)
        contenu = json.loads(serializers.serialize('json', list(qs)))
        _ecrire_classeur_systeme(wb, nom, contenu, _modele(nom))

    qs_ateliers = Atelier.objects.all().order_by('pk')
    qs_abos = Abonnement.objects.all().order_by('pk')
    qs_profils = Profil.objects.all().order_by('pk')
    qs_users = User.objects.all().order_by('pk')
    if atelier is not None:
        qs_ateliers = qs_ateliers.filter(pk=atelier.pk)
        qs_abos = qs_abos.filter(atelier=atelier)
        qs_profils = qs_profils.filter(atelier=atelier)
        qs_users = qs_users.filter(profil__atelier=atelier)
    feuille('Ateliers', ['ID', 'Nom', 'Ville', 'Créé le', 'Actif'],
            [[a.id, a.nom, getattr(a, 'ville', ''), a.date_creation, a.est_actif]
             for a in qs_ateliers])
    feuille('Abonnements', ['ID', 'Atelier', 'Plan', 'Statut', 'Début', 'Fin'],
            [[ab.id, str(ab.atelier), ab.plan.nom if ab.plan else '', ab.statut,
              ab.date_debut, ab.date_fin]
             for ab in qs_abos])
    feuille('Formules', ['ID', 'Nom', 'Prix mensuel', 'Durée'],
            [[p.id, p.nom, float(p.prix_mensuel), p.duree_libelle]
             for p in PlanAbonnement.objects.all().order_by('pk')])
    feuille('Profils', ['ID', 'Utilisateur', 'Rôle', 'Atelier', 'Fondateur'],
            [[p.id, p.user.username, p.role,
              p.atelier.nom if p.atelier else '', p.est_fondateur]
             for p in qs_profils])
    feuille('Comptes', ['ID', 'Nom d’utilisateur', 'Courriel', 'Prénom', 'Nom', 'Actif'],
            [[u.id, u.username, u.email, u.first_name, u.last_name, u.is_active]
             for u in qs_users])

    if atelier is not None:
        slug = re.sub(r'[^A-Za-z0-9_-]+', '-', atelier.nom)[:30].strip('-') or 'boutique'
        nom = f"sauvegarde_boutique_{slug}_{datetime.now():%Y%m%d-%H%M%S}.xlsx"
    else:
        nom = f"sauvegarde_systeme_{datetime.now():%Y%m%d-%H%M%S}.xlsx"
    wb.save(dossier_systeme() / nom)
    return nom


def _ecrire_classeur_systeme(wb, nom, contenu, modele):
    ws = wb.create_sheet(FEUILLES.get(nom, nom)[:31])
    if not contenu:
        ws.append(['(aucun enregistrement)'])
        return
    champs = []
    for objet in contenu:
        for cle in objet['fields']:
            if cle not in champs:
                champs.append(cle)
    ws.append(['Identifiant'] + [c.capitalize().replace('_', ' ') for c in champs])
    cache = {}
    for objet in contenu:
        ligne = [objet['pk']]
        for cle in champs:
            field = None
            try:
                field = modele._meta.get_field(cle)
            except Exception:
                field = None
            ligne.append(_valeur_lisible(objet['fields'].get(cle), field, cache))
        ws.append(ligne)


def lister_sauvegardes_systeme():
    resultats = []
    for chemin in dossier_systeme().glob('sauvegarde_systeme_*.xlsx'):
        motif = MOTIF_NOM_SYSTEME.match(chemin.name)
        if not motif:
            continue
        resultats.append({
            'nom': chemin.name,
            'date': datetime.strptime(motif.group(1), '%Y%m%d-%H%M%S'),
            'taille': chemin.stat().st_size,
        })
    resultats.sort(key=lambda e: e['date'], reverse=True)
    return resultats


def chemin_sauvegarde_systeme(nom):
    if not MOTIF_NOM_SYSTEME.match(nom or ''):
        return None
    chemin = dossier_systeme() / nom
    return chemin if chemin.is_file() else None


def supprimer_sauvegarde_systeme(nom):
    chemin = chemin_sauvegarde_systeme(nom)
    if chemin is None:
        return False
    chemin.unlink()
    return True
